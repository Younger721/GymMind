from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT = Path('output')
OUT.joinpath('docx').mkdir(parents=True, exist_ok=True)
OUT.joinpath('pdf').mkdir(parents=True, exist_ok=True)

sections = [
('健身房服务总览', [
 'GymMind 健身房为会员提供会员管理、教练服务、课程预约、训练计划、训练记录、动作库、营养建议和 AI 健身助手。会员可以查看课程、预约课程、记录训练，并针对个人知识库提问。',
 '系统中的 AI 负责理解、分析和推荐；会员状态、课程名额、训练容量、BMI、BMR、TDEE 和营养素等确定性数据由后端计算和校验。'
]),
('会员与教练服务', [
 '会员档案包括健身目标、健身经验、训练偏好、当前套餐、私教教练、最近签到和会员状态。教练可以管理会员、训练计划、训练记录和预约。',
 '教练资料包括擅长方向、简介、从业年限、证书、可预约时间和当前课程。会员选择教练后，需要提交申请；教练接受后关系变为有效。'
]),
('课程预约规则', [
 '课程类型包括团体课、私教课、体验课和训练课程。每门课程包含教练、课程说明、人数上限、时长、开始时间、结束时间、地点和状态。',
 '预约状态包括待处理、已确认、已取消、已完成和未到场。会员可在有余量时预约课程，并在规则允许的时间内取消预约。签到记录用于计算会员活跃度和运营分析。'
]),
('训练与动作安全', [
 '动作库按胸部、背部、肩部、手臂、腿部、臀部、核心、有氧和拉伸分类。每个动作记录目标肌群、难度、所需器械、步骤、常见错误和安全注意事项。',
 '新手进行力量训练时，应先完成热身，使用能保持动作稳定的重量，优先学习正确动作路径。出现锐痛、眩晕、胸闷或异常呼吸时应立即停止，并寻求教练或专业医疗人员帮助。AI 建议不能替代诊断。',
 '第三方教学视频只保存标题、原始 URL、平台、作者、缩略图和动作关联信息，不未经授权下载或重新托管。健身房拥有版权的内部视频才可以上传对象存储。'
]),
('训练计划与记录', [
 '训练计划包含训练目标、起止日期、状态和说明。计划项目包括动作、训练日、组数、次数、重量、组间休息和备注。',
 '训练记录包括训练日期、时长、消耗热量、动作、每组重量、次数、完成状态和备注。系统可统计训练次数、每周频率、总训练量、动作次数、重量变化、训练时长和连续训练天数。',
 'AI 生成计划后，后端必须校验动作 ID、组数、次数、休息时间、日期、权限和业务规则；AI 不直接写入数据库。'
]),
('营养建议边界', [
 '营养档案可以记录目标热量、蛋白质、碳水化合物、脂肪目标、饮食偏好和忌口。AI 可以生成一般性饮食建议，但热量和营养素数值由后端确定性计算。',
 '涉及疾病、孕期、进食障碍、药物相互作用或其他医疗风险的问题，应提示会员咨询医生或注册营养师，不能只依据 AI 回答做决定。'
]),
('RAG 知识库与回答规则', [
 '知识库支持 PDF、Word、Markdown、TXT、网页文章、健身房内部资料和教练上传资料。文档经过解析、清洗、分块、Embedding 后进入检索系统。',
 'RAG 查询流程是：问题意图判断、用户权限过滤、向量检索、全文检索、融合排序、重排、上下文压缩和回答生成。向量检索使用 Milvus，全文检索和元数据过滤使用 Elasticsearch。',
 '回答应优先引用实际召回的知识来源。资料不足时，系统必须明确说明“当前知识库没有足够依据”，不能编造来源、课程、价格或服务。所有私有知识检索必须按当前用户和当前健身房隔离。'
]),
('可用于测试的问题', [
 '健身房有哪些课程类型？',
 '课程预约状态有哪些？',
 '新手做力量训练要注意什么？',
 '第三方健身视频可以直接上传到健身房吗？',
 'AI 生成的训练计划为什么还要后端校验？',
 'RAG 找不到依据时应该怎么回答？',
 '会员档案中会记录哪些训练相关信息？',
 '哪些健康问题需要建议会员咨询专业人士？'
]),
]

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def build_docx(path):
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
    styles = doc.styles
    styles['Normal'].font.name = 'Microsoft YaHei'; styles['Normal'].font.size = Pt(10.5)
    styles['Title'].font.name = 'Microsoft YaHei'; styles['Title'].font.size = Pt(22); styles['Title'].font.bold = True; styles['Title'].font.color.rgb = RGBColor(0,0,0)
    for s in ['Heading 1','Heading 2']:
        styles[s].font.name = 'Microsoft YaHei'; styles[s].font.color.rgb = RGBColor(0,0,0)
    p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('GymMind 健身房 RAG 测试知识资料')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('用于知识库上传、检索、引用和无依据回答测试'); r.italic=True; r.font.color.rgb=RGBColor(90,90,90)
    doc.add_paragraph('本文档整理健身房服务、课程、训练安全、营养边界和 RAG 回答规则。内容采用短段落和明确术语，适合作为个人知识库测试资料。')
    for title, paras in sections:
        doc.add_heading(title, level=1)
        for t in paras:
            p=doc.add_paragraph(t); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.15
    doc.save(path)

def footer(canvas, doc):
    canvas.saveState(); canvas.setFont('Helvetica', 8); canvas.setFillColor(colors.grey); canvas.drawCentredString(A4[0]/2, 24, f'GymMind RAG 测试知识资料  |  {doc.page}'); canvas.restoreState()

def build_pdf(path):
    pdfmetrics.registerFont(TTFont('SimSun', r'C:\Windows\Fonts\Deng.ttf'))
    styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name='CNTitle', parent=styles['Title'], fontName='SimSun', fontSize=20, leading=26, alignment=TA_CENTER, textColor=colors.black, spaceAfter=12)); styles.add(ParagraphStyle(name='CNH1', parent=styles['Heading1'], fontName='SimSun', fontSize=14, leading=19, textColor=colors.HexColor('#183B56'), spaceBefore=12, spaceAfter=6)); styles.add(ParagraphStyle(name='CNBody', parent=styles['BodyText'], fontName='SimSun', fontSize=10, leading=16, spaceAfter=7)); styles.add(ParagraphStyle(name='CNSub', parent=styles['BodyText'], fontName='SimSun', fontSize=9, leading=13, alignment=TA_CENTER, textColor=colors.grey, spaceAfter=14))
    doc=SimpleDocTemplate(str(path), pagesize=A4, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=42)
    story=[Paragraph('GymMind 健身房 RAG 测试知识资料', styles['CNTitle']), Paragraph('用于知识库上传、检索、引用和无依据回答测试', styles['CNSub']), Paragraph('本文档整理健身房服务、课程、训练安全、营养边界和 RAG 回答规则。内容采用短段落和明确术语，适合作为个人知识库测试资料。', styles['CNBody'])]
    for title, paras in sections:
        story.append(Paragraph(title, styles['CNH1']))
        for t in paras: story.append(Paragraph(t, styles['CNBody']))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)

build_docx(OUT/'docx'/'GymMind_RAG测试知识资料.docx')
build_pdf(OUT/'pdf'/'GymMind_RAG测试知识资料.pdf')
