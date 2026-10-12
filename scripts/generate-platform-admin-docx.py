# -*- coding: utf-8 -*-
"""生成平台管理员应发布的知识库 Word 样例。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "rag-samples" / "平台管理员门店运营与合规发布手册.docx"


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)

    title = doc.add_heading("GymMind 平台门店运营与合规发布手册", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run("文档性质：").bold = True
    p.add_run("平台级知识库文件，由平台管理员上传并发布。")
    doc.add_paragraph("适用受众：各门店店管理员（GYM_ADMIN）、区域运营。")
    doc.add_paragraph("版本：V1.0　生效日期：2026-04-01　编制：GymMind 平台运营中心")
    doc.add_paragraph()

    doc.add_heading("一、发布说明", level=1)
    doc.add_paragraph(
        "本手册规定门店在 GymMind 平台上架课程、维护会员数据、执行安全巡检及向总部报送材料的统一要求。"
        "各门店店管理员应在系统「知识库」中确保员工可检索到与本手册一致的现行版本；"
        "修订后须重新上传并发布，旧版本由平台自动切换索引 generation，无需手工删库。"
    )

    doc.add_heading("二、知识库发布规范（平台管理员操作）", level=1)
    items = [
        "平台统一规范、器械标准、安全红线类文档：可见范围选择「平台」，受众角色勾选「店管理员（GYM_ADMIN）」。",
        "仅本店适用的促销、价目、排班说明：可见范围选择「本租户」，由店管理员上传并指定受众（会员/教练/管理员）。",
        "上传 Word（.docx）或 Markdown 后，等待状态变为 READY 再点击「发布」；未发布文档不参与 AI 检索。",
        "文件名建议包含主题与版本，例如「平台器械季检规范-V1.docx」，便于审计与 RAG 引用溯源。",
        "含个人手机号、银行尾号、内部密码的内容不得写入平台公开知识库，应使用租户私密文档或线下制度。",
    ]
    for i, text in enumerate(items, 1):
        doc.add_paragraph(f"{i}. {text}", style="List Number")

    doc.add_heading("三、门店开业与合规检查清单", level=1)
    doc.add_paragraph("新店开业或重装开业前 7 日，店管理员须在平台提交以下材料扫描件：")
    checklist = [
        "消防与疏散示意图已在训练区、更衣区张贴，应急灯与灭火器在有效期内。",
        "有氧器械（跑步机、椭圆机）已完成首次接地与急停绳检测，检测报告由店管理员签字存档。",
        "史密斯架、深蹲架安全销齐全，每日开班前点检表已配置到前台值班手册。",
        "AED（如有配置）电量与电极片在有效期内，当班至少 2 名员工完成心肺复苏培训。",
        "团体课程教练资质与保险在有效期内，已在 GymMind 后台绑定教练账号。",
    ]
    for text in checklist:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("四、器械与场地巡检（统一标准）", level=1)
    doc.add_paragraph(
        "【PLAT-OPS-001】跑步机与椭圆机每季度必须检查：急停绳、跑带磨损、电源接地；"
        "检查记录由店管理员签字，保存不少于 24 个月，并可在总部抽查时 24 小时内提供电子版。"
    )
    doc.add_paragraph(
        "【PLAT-OPS-002】史密斯架与深蹲架的安全销须每日开班前点检；"
        "缺失或变形销钉的器械应立即贴红色封条停用，并在 GymMind 工单中报修，不得带病开放。"
    )
    doc.add_paragraph(
        "【PLAT-OPS-003】自由重量区高峰时段应保证至少一名持有急救资质的当班员工；"
        "会员出现胸痛、眩晕或疑似中暑时，按训练区紧急流程报警并保护现场。"
    )

    doc.add_heading("五、课程上架与会员服务", level=1)
    doc.add_paragraph(
        "团课与私教课上架前，店管理员须确认：课程描述不含虚假承诺、价格与前台价目一致、"
        "取消规则与会员服务规则一致（团课开课前 2 小时可免费取消，否则记爽约）。"
    )
    doc.add_paragraph(
        "平台禁止门店在知识库或课程详情中使用「保证减重 X 公斤」「医疗疗效」等表述；"
        "涉及营养建议时应提示「不构成医疗建议，特殊人群请咨询医生」。"
    )

    doc.add_heading("六、数据报送与对账", level=1)
    doc.add_paragraph(
        "【PLAT-OPS-004】各门店须于每月 5 日前在平台财务模块上传上月银行回单或对账截图；"
        "与总部虚拟子账户相关的疑问，通过平台工单联系财务，不得在会员群公开讨论内部账户信息。"
    )
    doc.add_paragraph(
        "会员个人信息导出仅限店管理员且需「会员:导出」权限；导出文件 24 小时内删除，禁止外传。"
    )

    doc.add_heading("七、违规处理", level=1)
    doc.add_paragraph(
        "未按本手册发布现行安全规范、或未执行【PLAT-OPS-001】季检留档的门店，"
        "平台运营可暂停该门店新课程上架权限，直至补交材料并通过复核。"
    )
    doc.add_paragraph(
        "发生训练区重伤或监管抽查不合格事件，店管理员须在 2 小时内通过平台上报事件摘要，"
        "并在 72 小时内提交完整《训练区事件记录表》扫描件。"
    )

    doc.add_heading("八、修订记录", level=1)
    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "版本"
    hdr[1].text = "日期"
    hdr[2].text = "修订人"
    hdr[3].text = "摘要"
    row = table.rows[1].cells
    row[0].text = "V1.0"
    row[1].text = "2026-04-01"
    row[2].text = "平台运营"
    row[3].text = "首版发布，供各门店 GYM_ADMIN 检索与执行"

    doc.add_paragraph()
    doc.add_paragraph("—— 文档结束 ——").alignment = WD_ALIGN_PARAGRAPH.CENTER

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)
    print("size", OUT.stat().st_size)


if __name__ == "__main__":
    main()
