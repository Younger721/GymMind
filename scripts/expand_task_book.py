# -*- coding: utf-8 -*-
"""扩充任务书「主要任务与具体要求」栏目内容至1500字以上。

基于精简版任务书生成扩充版：保留原表格结构与段落格式，
仅替换第5行单元格中的正文段落。
"""
import copy
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document

SRC = r"D:\java\project\full-project\GymMind\本科毕业论文任务书-黄洋-精简版.docx"
DST = r"D:\java\project\full-project\GymMind\本科毕业论文任务书-黄洋-扩充版.docx"

# 新正文段落：(类型, 文本)，h 为小标题（沿用原小标题格式），b 为正文段落
PARAGRAPHS = [
    ("h", "（一）主要任务"),
    (
        "b",
        "本课题面向健身场馆数字化运营与会员科学健身需求，设计并实现基于B/S架构的"
        "智能健身信息管理与知识问答系统GymMind。当前健身场馆普遍存在会员信息分散、"
        "课程与器材管理效率不高、健身知识获取渠道零散等问题，管理员、教练与会员之间"
        "缺乏统一的信息化协作平台；同时，通用大语言模型在垂直领域问答中容易产生与场馆"
        "实际业务规则不符的内容，难以直接满足健身知识服务的准确性要求。因此，如何在"
        "健身场景下构建可信、可解释的知识问答能力，是本课题需要解决的核心问题。",
    ),
    (
        "b",
        "系统后端采用Spring Boot，前端采用Vue 3，业务数据存储于MySQL并结合Redis"
        "提升热点数据访问性能；在健身知识服务环节引入检索增强生成（RAG）技术，通过"
        "Elasticsearch实现关键词全文检索、Milvus实现向量语义检索，并对双路召回结果"
        "进行融合排序，在提高问答准确性的同时给出可追溯的引用依据，降低生成内容出现"
        "“幻觉”的风险。",
    ),
    (
        "b",
        "主要工作包括：查阅中英文文献并完成需求分析；完成系统总体设计、数据库设计"
        "与接口设计；完成前后端功能开发与RAG链路集成；执行功能测试、接口测试、安全"
        "测试与性能验证；撰写符合学术规范的毕业论文并整理各阶段过程材料。",
    ),
    ("h", "（二）具体要求"),
    (
        "b",
        "1. 文献调研与需求分析：查阅近五年中英文期刊与学术资料，梳理健身房信息管理"
        "系统、前后端分离Web应用、RAG与混合检索、Web软件测试等方向的研究现状与发展"
        "趋势，总结可借鉴的技术方案；在此基础上围绕管理员、教练、会员三类角色建立用例"
        "模型，明确各角色的功能需求、数据需求以及性能、安全等非功能需求，形成结构完整、"
        "术语规范的需求说明文档。",
    ),
    (
        "b",
        "2. 系统设计与数据库设计：采用B/S结构、前后端分离与分层设计思想，完成用例图、"
        "总体架构图、功能模块图、E-R图及注册登录、知识问答等关键业务流程图；数据库设计"
        "应满足完整性、一致性与可扩展性要求，表结构、索引与约束定义清晰，并为后续功能"
        "扩展预留空间；RAG子系统需说明知识文档切分、向量化、双路召回与融合排序的完整"
        "方案，并考虑知识库按用户隔离的检索策略；接口采用RESTful风格并统一响应格式，"
        "系统配置与密钥不得硬编码。",
    ),
    (
        "b",
        "3. 系统功能实现：实现用户注册登录、JWT鉴权与个人健身档案维护；实现健身知识"
        "文档的上传、解析、切分、入库与检索；实现基于RAG的多轮智能问答，支持上下文"
        "连续对话并展示引用依据；实现管理端对用户、知识库与运营数据的管理。后端需完善"
        "输入参数校验、权限控制、异常处理与日志记录，保证核心业务可追踪、可维护。",
    ),
    (
        "b",
        "4. 技术实现：后端使用Java 17与Spring Boot 3.x，采用Spring Security实现认证"
        "与授权，数据层使用MySQL 8.0与Redis；检索层部署Elasticsearch与Milvus，知识"
        "文档等附件存储使用MinIO；前端使用Vue 3、TypeScript、Vite、Pinia与Element Plus"
        "完成页面开发与状态管理；Elasticsearch、Milvus等中间件通过Docker Compose统一"
        "部署，便于开发与部署环境复现。",
    ),
    (
        "b",
        "5. 系统测试与验证：制定测试计划并执行功能测试、接口测试与主要页面测试，覆盖"
        "典型用户操作路径与异常输入场景，记录缺陷及修复情况；对RAG问答链路进行抽样"
        "验证，检验检索命中率与回答质量；开展鉴权、越权访问等安全测试，并给出主要接口"
        "响应时间与基本并发抽测结果，据此对系统进行优化改进。",
    ),
    (
        "b",
        "6. 其他要求：参考文献任务书不少于10篇、开题报告不少于12篇、论文不少于15篇，"
        "其中外文文献均不少于3篇，三份文件的参考文献应有区分度、不得雷同；本栏主要任务"
        "与具体要求不少于1500字；论文正文（不含附录）不少于15000字，应体现完整的软件"
        "工程开发过程，不宜大段粘贴项目源码，关键代码可置于附录；各阶段工作须按进度安排"
        "的时间节点完成，遵守学术规范并完成查重检测。",
    ),
]


def main() -> None:
    # 复制精简版生成扩充版，保留全部原有格式
    shutil.copyfile(SRC, DST)
    doc = Document(DST)
    cell = doc.tables[0].rows[4].cells[0]
    paras = list(cell.paragraphs)
    header, heading_tpl, body_tpl = paras[0], paras[1], paras[2]

    # 删除表头以外的原有正文段落
    for p in paras[1:]:
        p._element.getparent().remove(p._element)

    # 以模板段落格式生成新段落：深拷贝XML保留字体、对齐与行距
    for kind, text in PARAGRAPHS:
        tpl = heading_tpl if kind == "h" else body_tpl
        el = copy.deepcopy(tpl._element)
        cell._tc.append(el)
        p = cell.paragraphs[-1]
        # 仅保留第一个run并写入新文本，避免残留旧内容
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
        p.runs[0].text = text

    doc.save(DST)

    # 重新打开校验内容与字符数
    doc2 = Document(DST)
    cell2 = doc2.tables[0].rows[4].cells[0]
    body_text = "".join(p.text for p in cell2.paragraphs[1:])
    count = len("".join(body_text.split()))
    print("正文段落字符数（不含空白）:", count)
    for i, p in enumerate(cell2.paragraphs):
        print(f"[{i}]", p.text[:40])
    assert count >= 1500, f"字符数不足1500，当前仅{count}"
    assert cell2.paragraphs[0].text.startswith("论文（设计）的主要任务"), "表头段落丢失"
    print("校验通过")


if __name__ == "__main__":
    main()
