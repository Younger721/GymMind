#!/usr/bin/env python3
"""合并门店演示 MD 为单个 docx，删除旧 MD 并重新上传（需本地后端 8080 + gymadmin）。"""
from __future__ import annotations

import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
DEMO_DIR = ROOT / "backend" / "src" / "main" / "resources" / "demo-knowledge"
OUT_DOCX = DEMO_DIR / "健身房知识汇总手册.docx"
BASE = "http://localhost:8080"
USERNAME = "gymadmin"
PASSWORD = "123456"

SECTIONS: list[tuple[str, str]] = [
    ("常见健身问题.md", "faq.md"),
    ("健身房规章制度.md", "gym-rules.md"),
    ("饮食基础知识.md", "nutrition-basics.md"),
    ("团体课程介绍.md", "course-intro.md"),
    ("史密斯机使用说明.md", "equipment-guide.md"),
    ("训练基础知识.md", "training-basics.md"),
    ("深蹲动作规范.md", "action-standard.md"),
]

TARGET_MD_NAMES = {title for title, _ in SECTIONS}


def api(method: str, path: str, token: str | None = None, data: bytes | None = None, headers: dict | None = None):
    url = BASE + path
    h = dict(headers or {})
    if token:
        h["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            body = resp.read()
            if not body:
                return None
            return json.loads(body.decode("utf-8"))
    except urllib.error.HTTPError as ex:
        err = ex.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"{method} {path} -> {ex.code}: {err}") from ex


def login() -> str:
    payload = json.dumps({"username": USERNAME, "password": PASSWORD}).encode("utf-8")
    r = api("POST", "/api/v1/auth/login", data=payload, headers={"Content-Type": "application/json"})
    return r["data"]["accessToken"]


def append_md_to_doc(doc: Document, md_text: str) -> None:
    for line in md_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped.lstrip("#").strip()
            doc.add_heading(title, level=min(level, 3))
            continue
        if stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
            continue
        doc.add_paragraph(stripped)


def build_docx() -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    doc.add_heading("健身房知识汇总手册", 0)
    doc.add_paragraph("本文档由门店原有多份 Markdown 资料合并生成，供 AI 检索与员工查阅。")
    for section_title, resource_name in SECTIONS:
        md_path = DEMO_DIR / resource_name
        text = md_path.read_text(encoding="utf-8")
        doc.add_page_break()
        doc.add_heading(section_title.replace(".md", ""), level=1)
        append_md_to_doc(doc, text)
    doc.save(OUT_DOCX)
    return OUT_DOCX.read_bytes()


def upload_docx(token: str, content: bytes) -> int:
    boundary = "----GymMindBoundary7MA4YWxk"
    filename = OUT_DOCX.name
    parts: list[bytes] = []
    disposition = (
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        f"Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n"
    ).encode("utf-8")
    parts.append(f"--{boundary}\r\n".encode() + disposition + content + b"\r\n")
    parts.append(f"--{boundary}--\r\n".encode())
    body = b"".join(parts)
    params = urllib.parse.urlencode(
        [
            ("visibility", "TENANT"),
            ("category", "GENERAL"),
            ("knowledgeType", "GENERAL_NOTICE"),
        ]
    )
    for role in ("GYM_ADMIN", "COACH", "MEMBER"):
        params += "&" + urllib.parse.urlencode({"audienceRoles": role})
    path = f"/api/v1/knowledge/documents?{params}"
    h = {
        "Authorization": f"Bearer {token}",
        "Content-Type": f"multipart/form-data; boundary={boundary}",
    }
    req = urllib.request.Request(BASE + path, data=body, headers=h, method="POST")
    with urllib.request.urlopen(req, timeout=120) as resp:
        r = json.loads(resp.read().decode("utf-8"))
    return int(r["data"]["id"])


def wait_ready(token: str, doc_id: int, timeout_sec: int = 180) -> None:
    deadline = time.time() + timeout_sec
    while time.time() < deadline:
        row = api("GET", f"/api/v1/knowledge/documents/{doc_id}", token=token)["data"]
        status = row["status"]
        if status == "READY":
            return
        if status == "FAILED":
            raise RuntimeError(f"文档 {doc_id} 解析失败")
        time.sleep(3)
    raise RuntimeError(f"等待文档 {doc_id} READY 超时")


def publish(token: str, doc_id: int) -> None:
    api("PATCH", f"/api/v1/knowledge/documents/{doc_id}/publication?published=true", token=token)


def main() -> int:
    print("生成 Word…", OUT_DOCX)
    docx_bytes = build_docx()
    token = login()
    docs = api("GET", "/api/v1/knowledge/documents", token=token)["data"]
    to_delete = [d for d in docs if d.get("fileName") in TARGET_MD_NAMES]
    if not to_delete:
        print("未找到待删除的 MD（可能已处理过）")
    for d in sorted(to_delete, key=lambda x: x["id"]):
        print(f"删除 id={d['id']} {d['fileName']}")
        api("DELETE", f"/api/v1/knowledge/documents/{d['id']}", token=token)
    print("上传汇总 docx…")
    new_id = upload_docx(token, docx_bytes)
    print(f"新文档 id={new_id}，等待解析…")
    wait_ready(token, new_id)
    publish(token, new_id)
    print(f"完成：{OUT_DOCX.name}（id={new_id}）已发布")
    return 0


if __name__ == "__main__":
    sys.exit(main())
